"""01 - Extraction des résultats électoraux publiés par le Ministère de l'Intérieur (xlsx).
Filtre le département de l'Isère (38) et passe au format long (une ligne par unité x candidat)."""
import openpyxl, csv, sys

DEP = "38"

def extract(path, out, fixed, block_start, block_size, block_idx, block_names):
    """fixed : dict nom_colonne -> index dans la ligne source"""
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    next(rows)  # en-tête
    n = 0
    with open(out, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(list(fixed) + block_names)
        for r in rows:
            if r[0] is None or str(r[0]).zfill(2) != DEP:
                continue
            base = [r[i] for i in fixed.values()]
            i = block_start
            while i + block_size <= len(r) and r[i] is not None:
                blk = r[i:i + block_size]
                w.writerow(base + [blk[j] for j in block_idx])
                n += 1
                i += block_size
    print(out, n, "lignes", flush=True)

src = sys.argv[1]
which = sys.argv[2]
if which == "pres2022":
    # bloc candidat (7 col.) : panneau, sexe, nom, prénom, voix, %ins, %exp
    extract(f"{src}/pres2022_t1_burvot.xlsx", "data_raw/pres2022_t1_isere_bureaux.csv",
            {"com": 4, "libelle": 5, "bureau": 6, "inscrits": 7, "votants": 10, "blancs": 12, "nuls": 15, "exprimes": 18},
            21, 7, [2, 3, 4], ["nom", "prenom", "voix"])
elif which == "legi2022":
    # bloc candidat (8 col.) : panneau, sexe, nom, prénom, nuance, voix, %ins, %exp
    extract(f"{src}/legi2022_t1_subcom.xlsx", "data_raw/legi2022_t1_isere_communes.csv",
            {"circo": 2, "com": 4, "libelle": 5, "inscrits": 7, "votants": 10, "blancs": 12, "nuls": 15, "exprimes": 18},
            21, 8, [2, 3, 4, 5], ["nom", "prenom", "nuance", "voix"])
elif which == "legi2024":
    # bloc candidat (9 col.) : panneau, nuance, nom, prénom, sexe, voix, %ins, %exp, élu
    extract(f"{src}/legi2024_communes.xlsx", "data_raw/legi2024_t1_isere_communes.csv",
            {"com": 2, "libelle": 3, "inscrits": 4, "votants": 5, "blancs": 12, "nuls": 15, "exprimes": 9},
            18, 9, [1, 2, 3, 5], ["nuance", "nom", "prenom", "voix"])
